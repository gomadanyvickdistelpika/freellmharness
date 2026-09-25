"""Attachments, the folder fence, and scheduling.

The two tests that earn their keep here are the fence — where every escape
attempt is tried against a real symlink and a real traversal — and the
catch-up logic, where the failure mode is waking up to fourteen identical job
searches.
"""

from __future__ import annotations

import asyncio
import io
import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from harness import check, section  # noqa: E402


# ---------------------------------------------------------------------------
# 1. Extraction
# ---------------------------------------------------------------------------

def _extraction() -> None:
    from aegis import files

    result = files.extract("notes.txt", b"hello world")
    check("plain text extracts", result.text == "hello world" and result.kind == "text")

    result = files.extract("notes.md", "café — naïve".encode("utf-8"))
    check("utf-8 survives", "café" in result.text, result.text)

    result = files.extract("legacy.txt", "café".encode("cp1252"))
    check("a cp1252 file is decoded, not mangled into an exception",
          "caf" in result.text, repr(result.text))

    result = files.extract("data.csv", b"name,role\nAlex,engineer\nX,Y\n")
    check("csv becomes readable rows",
          "Alex | engineer" in result.text, result.text[:60])
    check("csv reports its shape", "3 rows" in result.text, result.text[:30])

    result = files.extract("cfg.json", b'{"b":2,"a":1}')
    check("json is pretty-printed", '"a": 1' in result.text, result.text)
    result = files.extract("broken.json", b'{not json')
    check("malformed json falls back to raw text", "not json" in result.text)

    result = files.extract("photo.png", b"\x89PNG\r\n\x1a\n")
    check("an image is flagged for the model, not extracted",
          result.kind == "image" and not result.text)

    result = files.extract("thing.bin", b"\x00\x01\x02\x03" * 100)
    check("binary is refused with a reason",
          result.kind == "binary" and "cannot read" in result.note, result.note)

    result = files.extract("script.py", b"print('hi')")
    check("source code counts as text", result.text == "print('hi')")

    # docx, for real.
    try:
        import docx
        document = docx.Document()
        document.add_paragraph("Certificate renewal steps")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "vCenter"
        table.rows[0].cells[1].text = "8.0"
        buffer = io.BytesIO()
        document.save(buffer)
        result = files.extract("doc.docx", buffer.getvalue())
        check("docx paragraphs extract",
              "Certificate renewal steps" in result.text, result.text[:60])
        check("docx tables extract too", "vCenter | 8.0" in result.text,
              result.text[-60:])
    except ImportError:
        check("docx extraction", False, "python-docx is not installed")

    # xlsx, for real.
    try:
        import openpyxl
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Cases"
        sheet.append(["id", "status"])
        sheet.append([1, "closed"])
        buffer = io.BytesIO()
        book.save(buffer)
        result = files.extract("book.xlsx", buffer.getvalue())
        check("xlsx sheets are named in the output",
              "sheet: Cases" in result.text, result.text[:60])
        check("xlsx rows extract", "1 | closed" in result.text, result.text[-40:])
    except ImportError:
        check("xlsx extraction", False, "openpyxl is not installed")

    # A PDF with no text layer must say so rather than attaching nothing.
    try:
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        result = files.extract("scan.pdf", buffer.getvalue())
        check("a PDF with no text layer says so instead of failing silently",
              result.kind == "pdf" and "scan" in result.note.lower(), result.note)
        check("...and still reports the page count", result.pages == 1)
    except ImportError:
        check("pdf extraction", False, "pypdf is not installed")

    result = files.extract("junk.pdf", b"this is not a pdf at all")
    check("a corrupt PDF reports the failure",
          result.kind == "pdf" and bool(result.note), result.note[:60])


def _storage() -> None:
    from aegis import files

    chat_id = "chat-test"
    small = files.add(chat_id, "small.txt", b"a short note")
    check("a small file is stored", small["ok"] is True)
    check("...and marked to ride inline", small["inline"] is True)

    big_text = ("Lorem ipsum dolor sit amet. " * 2000).encode()
    big = files.add(chat_id, "big.txt", big_text)
    check("a big file is stored", big["ok"] is True)
    check("...and NOT inlined", big["inline"] is False,
          f"{big['chars']} chars vs limit {files.INLINE_LIMIT}")

    listed = files.for_chat(chat_id)
    check("both are listed against the chat", len(listed) == 2)
    check("the stored record hides the full text",
          "text" not in listed[0], str(listed[0].keys()))

    body = files.text_of(big["id"], 0, 100)
    check("text reads back in chunks",
          body["ok"] and len(body["text"]) == 100 and body["more"] is True)
    body = files.text_of(big["id"], 50, 100)
    check("an offset read starts where asked", body["offset"] == 50)
    check("an unknown id fails cleanly", files.text_of("nope")["ok"] is False)

    hits = files.search("Lorem", chat_id)
    check("attachments are searchable", any(h["attach_id"] == big["id"] for h in hits),
          str(len(hits)))
    check("search misses what is not there", files.search("zzzqqq", chat_id) == [])

    # The part that protects the context window.
    blocks, images = files.message_parts([small["id"], big["id"]])
    check("the small file's text is in the message",
          "a short note" in blocks, blocks[:80])
    check("the big file's text is NOT in the message",
          "Lorem ipsum" not in blocks)
    check("...and the big file is announced with its id",
          big["id"] in blocks and "files__search_attachments" in blocks,
          blocks[-160:])
    check("no images for text files", images == [])

    png = (b"\x89PNG\r\n\x1a\n" + b"\x00" * 40)
    shot = files.add(chat_id, "shot.png", png)
    blocks, images = files.message_parts([shot["id"]])
    check("an image becomes an image part",
          len(images) == 1 and images[0][0] == "image/png", str(images[:1]))

    # The real limit is 30 GB (v2.1); shrink it so the test does not need 30 GB.
    from aegis.config import settings as _settings
    _settings.set("max_upload_gb", 0.0001)
    over = files.add(chat_id, "huge.bin", b"x" * (files.max_upload() + 1))
    _settings.set("max_upload_gb", 30)
    check("an oversized upload is refused with the limit",
          over["ok"] is False and "limit" in over["error"], over.get("error", ""))

    check("an attachment can be deleted", files.delete(small["id"]) is True)
    check("deleting twice is not an error", files.delete(small["id"]) is False)

    nasty = files.add(chat_id, "../../evil .txt", b"x")
    check("a traversal filename is flattened on save",
          ".." not in nasty["name"] and "/" not in nasty["name"], nasty["name"])


# ---------------------------------------------------------------------------
# 2. The folder fence
# ---------------------------------------------------------------------------

def _fence() -> None:
    from aegis.config import settings
    from aegis.tools import files as ft

    settings.set("folders", [])
    # v2 gives the agent its own workspace by default; the fence rules below
    # are about *your* folders, so switch the workspace off for them.
    settings.set("workspace_enabled", False)
    check("with no folders connected there is no filesystem", not ft.folders())
    check("...and the only tools are the attachment ones",
          {t.raw_name for t in ft.tools()} == {"search_attachments", "read_attachment"},
          str({t.raw_name for t in ft.tools()}))

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp).resolve()
        (root / "work").mkdir()
        (root / "work" / "notes.md").write_text("# Notes\nvCenter cert expiry\n",
                                                encoding="utf-8")
        (root / "work" / "sub").mkdir()
        (root / "work" / "sub" / "deep.txt").write_text("deep value", encoding="utf-8")
        secret = root / "secret.txt"
        secret.write_text("SSH KEY", encoding="utf-8")

        result = ft.add_folder(str(root / "work"))
        check("a folder connects", result["ok"] is True)
        check("connecting a missing folder is refused",
              ft.add_folder(str(root / "nope"))["ok"] is False)
        check("connecting the same folder twice is refused",
              ft.add_folder(str(root / "work"))["ok"] is False)

        inside = ft.resolve(str(root / "work" / "notes.md"))
        check("a path inside resolves", inside.name == "notes.md")
        check("a relative path resolves against the first folder",
              ft.resolve("notes.md").name == "notes.md")
        check("a nested path resolves", ft.resolve("sub/deep.txt").name == "deep.txt")

        # Every escape route, tried for real.
        escapes = {
            "absolute path outside": str(secret),
            "dot-dot traversal": "../secret.txt",
            "nested dot-dot": "sub/../../secret.txt",
            "windows-style traversal": "..\\secret.txt",
        }
        for label, attempt in escapes.items():
            try:
                ft.resolve(attempt)
                check(f"{label} is refused", False, f"{attempt} got through")
            except ft.OutsideFence:
                check(f"{label} is refused", True)

        # A symlink is the one that beats naive string checks.
        try:
            link = root / "work" / "backdoor"
            link.symlink_to(secret)
            try:
                ft.resolve(str(link))
                check("a symlink pointing outside is refused", False,
                      "the symlink got through")
            except ft.OutsideFence:
                check("a symlink pointing outside is refused", True)
        except (OSError, NotImplementedError):
            check("a symlink pointing outside is refused", True,
                  "symlinks unavailable here; skipped")

        names = {t.raw_name for t in ft.tools()}
        check("connecting a folder adds the file tools",
              {"list", "read", "search", "write", "folders"} <= names, str(names))
        risks = {t.raw_name: t.risk for t in ft.tools()}
        check("reading is a read", risks["read"] == "read")
        check("listing is a read", risks["list"] == "read")
        check("searching is a read", risks["search"] == "read")
        check("writing is a write", risks["write"] == "write")

        listing = ft._list(path=str(root / "work"))
        check("listing shows files", "notes.md" in listing["text"])
        check("listing marks folders", "[dir]" in listing["text"])
        filtered = ft._list(path=str(root / "work"), pattern="*.md")
        check("a glob filters the listing",
              "notes.md" in filtered["text"] and "deep" not in filtered["text"])

        body = ft._read(path="notes.md")
        check("a file reads", "cert expiry" in body["text"])
        found = ft._search(query="cert", pattern="*.md")
        check("search finds the line", "notes.md" in found["text"], found["text"][:80])
        check("search reports a miss honestly",
              "No match" in ft._search(query="zzzqqq")["text"])

        written = ft._write(path="out/report.md", content="# Report\n")
        check("a write creates parent folders",
              (root / "work" / "out" / "report.md").exists(), written["text"])
        appended = ft._write(path="out/report.md", content="more\n", append=True)
        check("append adds rather than replaces",
              "more" in (root / "work" / "out" / "report.md").read_text()
              and "# Report" in (root / "work" / "out" / "report.md").read_text())

        escaped = ft._wrap(ft._write)
        outcome = asyncio.run(escaped({"path": str(secret), "content": "pwned"}))
        check("writing outside the fence is refused",
              outcome.get("error") is True and "Refused" in outcome["text"],
              outcome["text"][:70])
        check("...and the file outside is untouched",
              secret.read_text() == "SSH KEY")

        binary = root / "work" / "thing.bin"
        binary.write_bytes(b"\x00\x01" * 50)
        check("a binary file is not read as text",
              ft._read(path="thing.bin").get("error") is True)

        ft.remove_folder(str(root / "work"))
        check("a folder disconnects", not ft.folders())

    settings.set("folders", [])


# ---------------------------------------------------------------------------
# 3. Cron
# ---------------------------------------------------------------------------

def _cron() -> None:
    from aegis import scheduler as S

    cron = S.parse_cron("30 8 * * 1-5")
    check("a weekday cron parses", cron.hour == {8} and cron.minute == {30})
    check("a weekday range expands", cron.dow == {1, 2, 3, 4, 5})

    check("steps expand", sorted(S.parse_cron("*/15 * * * *").minute)
          == [0, 15, 30, 45])
    check("lists expand", S.parse_cron("0 9,17 * * *").hour == {9, 17})

    for bad in ("", "0 8 * *", "99 8 * * *", "0 25 * * *", "a b c d e"):
        try:
            S.parse_cron(bad)
            check(f"{bad!r} is rejected", False, "it parsed")
        except (ValueError, TypeError):
            check(f"{bad!r} is rejected", True)

    # Monday 2026-09-14 09:00; the next weekday 08:30 is Tuesday.
    monday = datetime(2026, 9, 14, 9, 0)
    upcoming = S.next_run("30 8 * * 1-5", monday)
    check("next run skips to the following weekday",
          upcoming == datetime(2026, 9, 15, 8, 30), str(upcoming))

    # Friday evening -> the next weekday run is Monday, not Saturday.
    friday = datetime(2026, 9, 18, 19, 0)
    upcoming = S.next_run("30 8 * * 1-5", friday)
    check("a weekday schedule skips the weekend",
          upcoming == datetime(2026, 9, 21, 8, 30), str(upcoming))

    earlier = S.previous_run("30 8 * * 1-5", datetime(2026, 9, 15, 9, 0))
    check("previous run finds this morning",
          earlier == datetime(2026, 9, 15, 8, 30), str(earlier))

    check("a weekday cron reads plainly",
          S.describe("0 8 * * 1-5") == "every weekday at 08:00",
          S.describe("0 8 * * 1-5"))
    check("a daily cron reads plainly",
          S.describe("30 7 * * *") == "every day at 07:30", S.describe("30 7 * * *"))
    check("an interval reads plainly",
          "15 minutes" in S.describe("*/15 * * * *"), S.describe("*/15 * * * *"))
    check("an invalid cron says so", "invalid" in S.describe("nonsense"))

    # Windows translation.
    check("a daily cron becomes a DAILY schtasks entry",
          S._schtasks_schedule("0 8 * * *") == ["/SC", "DAILY", "/ST", "08:00"],
          str(S._schtasks_schedule("0 8 * * *")))
    weekly = S._schtasks_schedule("30 7 * * 1-5")
    check("a weekday cron becomes WEEKLY with named days",
          weekly[:2] == ["/SC", "WEEKLY"] and "MON,TUE,WED,THU,FRI" in weekly,
          str(weekly))
    try:
        S._schtasks_schedule("*/15 * * * *")
        check("an every-N-minutes cron is refused for Windows", False)
    except ValueError as exc:
        check("an every-N-minutes cron is refused for Windows",
              "fixed time" in str(exc), str(exc))


# ---------------------------------------------------------------------------
# 4. Tasks and catch-up
# ---------------------------------------------------------------------------

def _tasks() -> None:
    from aegis import scheduler as S
    from aegis.config import settings

    settings.set("tasks", {})

    bad = S.create("", "do things", "0 8 * * *")
    check("a task needs a name", bad["ok"] is False)
    bad = S.create("Job search", "go", "not a cron")
    check("a task needs a valid cron", bad["ok"] is False, bad.get("error", ""))

    made = S.create("Daily job search", "Find remote VMware roles",
                    "0 8 * * 1-5", output_dir="", trusted=False)
    check("a task is created", made["ok"] is True)
    key = made["key"]
    check("the key is slugged from the name", key == "daily-job-search", key)

    task = S.get(key)
    check("it is stored", task is not None and task.name == "Daily job search")
    check("it is enabled by default", task.enabled is True)
    check("it is NOT trusted by default", task.trusted is False)

    shown = task.to_dict()
    check("the schedule is described for the UI",
          shown["schedule_text"] == "every weekday at 08:00", shown["schedule_text"])
    check("the next run is computed", shown["next_run"] > 0)

    again = S.create("Daily job search", "Something else", "0 9 * * *")
    check("a duplicate name gets a distinct key", again["key"] != key, again["key"])

    check("a task is deleted", S.delete(key) is True)
    check("deleting twice is not an error", S.delete(key) is False)
    S.delete(again["key"])
    settings.set("tasks", {})


def _catch_up() -> None:
    """Coming back from a week away must not produce seven job searches."""
    from aegis import scheduler as S
    from aegis.config import settings

    settings.set("tasks", {})
    made = S.create("Catchup probe", "do it", "0 8 * * *")
    key = made["key"]

    ran: list[str] = []

    async def fake_run(task_key: str, *, reason: str = "") -> dict[str, Any]:
        ran.append(f"{task_key}:{reason}")
        task = S.get(task_key)
        task.last_run = datetime.now().timestamp()
        S.save(task)
        return {"ok": True, "seconds": 0.0, "status": "ok", "file": ""}

    original = S.run_once
    S.run_once = fake_run
    try:
        scheduler = S.Scheduler()
        fired = asyncio.run(scheduler.catch_up())
        check("a missed run is caught up once on launch",
              fired == [key], str(fired))
        check("...and is labelled as a catch-up",
              ran and ran[0].endswith(":catch-up"), str(ran))

        ran.clear()
        fired = asyncio.run(scheduler.catch_up())
        check("a second launch does not re-run it", fired == [], str(fired))

        # A week of absence still yields exactly one run, not seven.
        task = S.get(key)
        task.last_run = (datetime.now().timestamp() - 7 * 86400)
        S.save(task)
        ran.clear()
        fired = asyncio.run(scheduler.catch_up())
        check("a week away produces one run, not one per missed day",
              len(fired) == 1, str(fired))

        # Too stale to bother with.
        task = S.get(key)
        task.last_run = 0.0
        task.catch_up = False
        S.save(task)
        check("catch-up can be switched off per task",
              asyncio.run(scheduler.catch_up()) == [])

        task = S.get(key)
        task.catch_up = True
        task.enabled = False
        task.last_run = 0.0
        S.save(task)
        check("a paused task is never caught up",
              asyncio.run(scheduler.catch_up()) == [])
    finally:
        S.run_once = original
        settings.set("tasks", {})


def _api() -> None:
    from fastapi.testclient import TestClient
    from aegis.server import app

    with TestClient(app) as client:
        r = client.get("/api/files")
        check("GET /api/files", r.status_code == 200 and "folders" in r.json())
        r = client.post("/api/files/folders", json={"path": "/definitely/not/here"})
        check("connecting a missing folder is refused by the API",
              r.status_code == 400)

        r = client.get("/api/tasks")
        check("GET /api/tasks", r.status_code == 200 and "scheduler" in r.json())

        r = client.post("/api/tasks", json={"name": "API probe",
                                            "prompt": "go", "cron": "bad"})
        check("a bad cron is refused by the API", r.status_code == 400)

        r = client.post("/api/tasks", json={"name": "API probe", "prompt": "go",
                                            "cron": "0 8 * * 1-5"})
        check("a task is created over the API", r.status_code == 200)
        key = r.json()["key"]

        r = client.get(f"/api/tasks/{key}/preview?cron=0%208%20*%20*%201-5")
        check("a cron can be previewed before saving",
              r.json()["ok"] and len(r.json()["next"]) >= 3, str(r.json())[:80])
        r = client.get("/api/tasks/_/preview?cron=nope")
        check("a bad cron preview explains itself", r.json()["ok"] is False)

        r = client.post("/api/tasks", json={"key": key, "enabled": False})
        check("a task can be paused over the API",
              r.json()["task"]["enabled"] is False)

        chat = client.post("/api/chats", json={}).json()["chat"]["id"]
        r = client.post(f"/api/chats/{chat}/attachments",
                        files={"uploads": ("note.txt", b"hello from the api",
                                           "text/plain")})
        check("a file uploads over the API",
              r.status_code == 200 and r.json()["attachments"][0]["ok"], r.text[:120])
        attach_id = r.json()["attachments"][0]["id"]

        r = client.get(f"/api/chats/{chat}/attachments")
        check("the chat lists its attachments", len(r.json()["attachments"]) == 1)
        r = client.get(f"/api/attachments/{attach_id}/text")
        check("the text reads back", "hello from the api" in r.json()["text"])
        r = client.get(f"/api/attachments/{attach_id}")
        check("the raw file downloads", r.content == b"hello from the api")
        check("an unknown attachment 404s",
              client.get("/api/attachments/nope").status_code == 404)

        check("an attachment deletes",
              client.delete(f"/api/attachments/{attach_id}").json()["ok"] is True)
        check("a task deletes", client.delete(f"/api/tasks/{key}").json()["ok"] is True)
        client.delete(f"/api/chats/{chat}")


# ---------------------------------------------------------------------------

def run_all() -> None:
    section("attachment extraction")
    _extraction()
    section("attachment storage and the context guard")
    _storage()
    section("the folder fence")
    _fence()
    section("cron")
    _cron()
    section("scheduled tasks")
    _tasks()
    section("missed runs")
    _catch_up()
    section("files and tasks API")
    _api()


if __name__ == "__main__":
    os.environ.setdefault("AEGIS_DATA_DIR",
                          tempfile.mkdtemp(prefix="aegis-files-test-"))
    from harness import report
    run_all()
    sys.exit(report())
